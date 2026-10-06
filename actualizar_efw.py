"""
actualizar_efw.py — reconstruye los datos del índice EFW desde el archivo maestro que el Fraser Institute
publica con cada informe (Economic Freedom of the World: <año> Annual Report):

  data/efw_panel_map.json     {año: {ISO3: {s, a1..a5}}} — hoja «EFW Panel Dataset» (1970 en adelante, quinquenal
                              hasta 1995 y anual desde 2000) + hoja «EFW Ratings 1950-1965» (sólo el índice general)
  data/efw_country_meta.json  {ISO3: {name, region}} de las jurisdicciones del último año
  data/consolidated.json      SÓLO los campos del índice (efw_avg, efw_a1..a5 = promedio 2000→último año,
                              efw_rank, efw_quartile, efw_n_years), region e income_class. Los indicadores de
                              bienestar (WDI, PIP, V-Dem, CPI…) vienen de otras fuentes y no se tocan.
  data/efw_componentes_bol.json  los componentes de Bolivia en la hoja oficial, con su valor de origen («· dato»)
                              y los puestos oficiales por área: con eso se explica en los textos QUÉ movió el índice.

    python actualizar_efw.py                         # usa (o descarga) el archivo del informe AÑO en data/
    python actualizar_efw.py ruta/archivo.xlsx
    python actualizar_efw.py ruta/archivo.xlsx --salida carpeta   # escribe en otra carpeta, para comparar

Controles (si uno falla, no se escribe nada):
  · el último año del panel tiene que coincidir con la hoja oficial (índice y cinco áreas de cada jurisdicción).
    ⚠️ En el archivo del informe 2026 el panel trae TODAS las áreas 0,01 por encima de la hoja oficial: 825 de 825
    en 2024 (Venezuela marca 0,01 en Moneda sana, que oficialmente es 0,00), y el índice general queda también
    0,01 arriba (Bolivia 5,48 en el panel, 5,47 en el informe). El script MIDE ese desfase en el último año y, si
    es parejo, lo descuenta en todo el panel; si no es parejo, se detiene. En el archivo 2025 el desfase es cero.
  · el informe que dice el nombre del archivo es el año del último dato + 2 (la regla que usan la página y los
    gráficos para citar la edición);
  · toda región se reconoce; todo nombre de la hoja 1950-1965 tiene su código ISO.
"""
import json
import math
import re
import statistics as st
import sys
import urllib.request
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent
DATA = RAIZ / 'data'
INFORME = 2026            # edición vigente: cambiarla acá cuando salga la siguiente
URL = 'https://efotw.org/sites/all/modules/custom/ftw_maps_pages/files/efotw-{}-master-index-data-for-researchers-iso.xlsx'
DESDE_PROMEDIO = 2000     # el promedio de consolidated.json va de este año al último (serie anual)

# Nombres de las regiones del Banco Mundial tal como los usan los gráficos (los del archivo 2025). El 2026 escribe
# «Latin America & Caribbean»; se lleva al nombre de siempre para no romper la traducción de los generadores.
REGIONES = {
    'East Asia & Pacific': 'East Asia & Pacific', 'Europe & Central Asia': 'Europe & Central Asia',
    'Latin America & the Caribbean': 'Latin America & the Caribbean', 'Latin America & Caribbean': 'Latin America & the Caribbean',
    'Middle East & North Africa': 'Middle East & North Africa', 'North America': 'North America',
    'South Asia': 'South Asia', 'Sub-Saharan Africa': 'Sub-Saharan Africa',
}
INGRESO = {'high income': 'High Income', 'upper-middle income': 'Upper-Middle Income',
           'lower-middle income': 'Lower-Middle Income', 'low income': 'Low Income'}
# La hoja 1950-1965 identifica por NOMBRE abreviado. Los que no coinciden con el nombre del panel se DECLARAN acá
# (hasta 2026 se cruzaba sólo por nombre exacto y quedaban fuera 15 o 16 países por año, Venezuela y Hong Kong entre ellos).
HIST_ISO = {
    'Central Afr. Rep.': 'CAF', 'Congo, Dem. R.': 'COD', 'Congo, Rep. Of': 'COG', "Cote d'Ivoire": 'CIV',
    'Czech Rep.': 'CZE', 'Dominican Rep.': 'DOM', 'Egypt': 'EGY', 'Hong Kong': 'HKG', 'Iran': 'IRN',
    'Korea, South': 'KOR', 'Pap. New Guinea': 'PNG', 'Russia': 'RUS', 'Syria': 'SYR', 'Trinidad & Tob.': 'TTO',
    'Turkey': 'TUR', 'Venezuela': 'VEN',
}
K = ('s', 'a1', 'a2', 'a3', 'a4', 'a5')


def alto(msg):
    raise SystemExit(f'✗ {msg} — no se escribió nada.')


def numero(v):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(x) else x


def filas_de(ws):
    return [r for r in ws.iter_rows(values_only=True)]


def leer(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    hojas = wb.sheetnames
    # ── panel encadenado (por nombre de columna: el orden cambió entre 2025 y 2026)
    fp = filas_de(wb['EFW Panel Dataset'])
    h = [str(x).strip() if x is not None else '' for x in fp[0]]
    iy, ii, ic = h.index('Year'), h.index('ISO_Code'), h.index('Countries')
    isum = next(j for j, x in enumerate(h) if x == 'Summary' or x.startswith('ECONOMIC FREEDOM'))
    iar = [next(j for j, x in enumerate(h) if x.startswith(f'Area {n}')) for n in range(1, 6)]
    panel, areas, nombre = {}, {}, {}
    for r in fp[1:]:
        if not r or r[iy] is None or r[ii] is None:
            continue
        rec = dict(zip(K, [numero(r[isum])] + [numero(r[j]) for j in iar]))
        iso = str(r[ii]).strip()
        # hay filas con áreas y sin índice general (países que todavía no se califican): no entran al panel,
        # pero sí al promedio de cada área en consolidated.json, como se calculó siempre
        areas.setdefault(int(r[iy]), {})[iso] = rec
        if rec['s'] is None:
            continue
        panel.setdefault(int(r[iy]), {})[iso] = rec
        nombre[iso] = str(r[ic]).strip()
    # ── hoja oficial (la del informe): índice, cinco áreas (Área 2 con el ajuste de género), región e ingreso
    hoja = hojas[0]
    fo = filas_de(wb[hoja])
    hi = next(i for i, r in enumerate(fo) if r and r[0] == 'Year')
    h = [str(x).strip() if x is not None else '' for x in fo[hi]]

    def col(cond):
        return next(j for j, x in enumerate(h) if cond(x))

    co = {'s': 3,
          'a1': col(lambda x: x.startswith('Area 1') and 'Rank' not in x),
          'a2': col(lambda x: x.startswith('Area 2') and 'With Gender' in x),
          'a3': col(lambda x: x.startswith('Area 3') and 'Rank' not in x),
          'a4': col(lambda x: x.startswith('Area 4') and 'Rank' not in x),
          'a5': col(lambda x: x.startswith('Area 5') and 'Rank' not in x)}
    jreg = col(lambda x: x.startswith('World Bank Region'))
    jing = col(lambda x: x.startswith('World Bank Current Income'))
    # rótulos de todas las columnas de la hoja oficial: la columna «data» que sigue a un componente es su valor de
    # origen (la inflación en %, el crecimiento del dinero…) y se rotula «<componente> · dato»
    rotulos, previo = {}, None
    for j, x in enumerate(h[:jreg]):
        if j < 3 or not x:
            continue
        if x == 'data':
            if previo:
                rotulos[j] = previo + ' · dato'
        else:
            previo = rotulos[j] = ' '.join(x.split())
    oficial, region, ingreso, componentes = {}, {}, {}, {}
    for r in fo[hi + 1:]:
        if not r or r[0] is None:
            continue
        y, iso = int(r[0]), str(r[1]).strip()
        oficial.setdefault(y, {})[iso] = {k: numero(r[j]) for k, j in co.items()}
        region.setdefault(y, {})[iso] = str(r[jreg]).strip() if r[jreg] else None
        ingreso.setdefault(y, {})[iso] = str(r[jing]).strip() if r[jing] else None
        if iso == 'BOL':
            fila = {}
            for j, rot in rotulos.items():
                v = numero(r[j])
                if v is None and r[j] not in (None, ''):
                    v = str(r[j]).strip()          # rangos como «26-43» (tasas marginales): quedan como texto
                if v is not None:
                    fila[rot] = round(v, 4) if isinstance(v, float) else v
            componentes[str(y)] = fila
    # ── 1950-1965: sólo el índice general, por nombre
    hist = {}
    for r in filas_de(wb[next(s for s in hojas if '1950' in s)])[1:]:
        if r and r[0] is not None and numero(r[2]) is not None:
            hist.setdefault(int(r[0]), {})[str(r[1]).strip()] = numero(r[2])
    return {'hojas': hojas, 'panel': panel, 'areas': areas, 'nombre': nombre, 'oficial': oficial, 'region': region,
            'ingreso': ingreso, 'hist': hist, 'componentes_bol': componentes}


def corregir_desfase(d):
    """Mide panel − oficial en el último año (tienen que coincidir) y, si el desfase es parejo, lo descuenta."""
    ult = max(d['panel'])
    pan, ofi = d['panel'][ult], d['oficial'].get(ult, {})
    if set(pan) != set(k for k, v in ofi.items() if v['s'] is not None):
        alto(f'el panel y la hoja oficial de {ult} no tienen las mismas jurisdicciones')
    difs = [round(pan[i][k] - ofi[i][k], 4) for i in pan for k in K[1:] if pan[i][k] is not None and ofi[i][k] is not None]
    desfase = st.median(difs)
    parejo = sum(1 for x in difs if abs(x - desfase) < 1e-4)
    if parejo != len(difs):
        alto(f'el desfase panel − oficial de {ult} no es parejo ({parejo} de {len(difs)} áreas en {desfase:+.4f})')
    if desfase:
        for y in d['areas']:             # el panel comparte los registros: se corrige una sola vez
            for rec in d['areas'][y].values():
                for k in K:
                    if rec[k] is not None:
                        rec[k] = max(0.0, rec[k] - desfase)
    # control: índice y áreas del último año, al centésimo, iguales a la hoja oficial
    mal = [(i, k, round(pan[i][k], 2), ofi[i][k]) for i in pan for k in K
           if pan[i][k] is not None and ofi[i][k] is not None and abs(round(pan[i][k], 2) - round(ofi[i][k], 2)) > 0.001]
    if mal:
        alto(f'{len(mal)} valores de {ult} no coinciden con la hoja oficial tras el ajuste, p. ej. {mal[:4]}')
    return desfase, len(difs), ult


def r2(x):
    return None if x is None else round(x + 0.0, 2)


def construir(d):
    ult = max(d['panel'])
    # ── panel: años en orden, jurisdicciones en orden alfabético, dos decimales (como siempre)
    iso_por_nombre = {v: k for k, v in d['nombre'].items()}
    out = {}
    sin_iso = set()
    for y in sorted(d['hist']):
        fila = {}
        for nom, s in d['hist'][y].items():
            iso = iso_por_nombre.get(nom) or HIST_ISO.get(nom)
            if not iso:
                sin_iso.add(nom)
                continue
            fila[iso] = {'s': r2(s), 'a1': None, 'a2': None, 'a3': None, 'a4': None, 'a5': None}
        out[str(y)] = dict(sorted(fila.items()))
    if sin_iso:
        alto(f'nombres de la hoja 1950-1965 sin código ISO: {sorted(sin_iso)}')
    for y in sorted(d['panel']):
        out[str(y)] = {i: {k: r2(v[k]) for k in K} for i, v in sorted(d['panel'][y].items())}
    # ── región e ingreso: los del último año
    reg = {}
    for i in d['panel'][ult]:
        r = d['region'][ult].get(i)
        if r not in REGIONES:
            alto(f'región no reconocida para {i}: {r!r}')
        reg[i] = REGIONES[r]
    ing = {i: INGRESO.get((d['ingreso'][ult].get(i) or '').lower(), d['ingreso'][ult].get(i) or 'NA') for i in d['panel'][ult]}
    # en el orden del archivo que ya existe (el diff muestra sólo lo que cambió); las nuevas, al final
    previo = DATA / 'efw_country_meta.json'
    orden = list(json.loads(previo.read_text(encoding='utf-8'))) if previo.exists() else []
    isos = [i for i in orden if i in d['panel'][ult]] + sorted(i for i in d['panel'][ult] if i not in orden)
    meta = {i: {'name': d['nombre'][i], 'region': reg[i]} for i in isos}
    # ── promedios 2000→último año, puesto y cuartil (cortes por cuantil, el sobrante en el cuartil 4)
    anios = [y for y in sorted(d['panel']) if y >= DESDE_PROMEDIO]
    prom = {}
    for i in d['panel'][ult]:
        # índice: los años calificados; cada área: todos los años con esa área (aunque falte el índice)
        recs = [d['areas'][y][i] for y in anios if i in d['areas'].get(y, {})]
        prom[i] = {k: st.fmean([x[k] for x in recs if x[k] is not None]) if any(x[k] is not None for x in recs) else None for k in K}
        prom[i]['n'] = sum(1 for x in recs if x['s'] is not None)
    orden = sorted(prom, key=lambda i: -prom[i]['s'])
    n = len(orden)
    vals = sorted(prom[i]['s'] for i in orden)

    def cuantil(p):          # interpolación lineal, como pandas
        x = p * (n - 1)
        lo = math.floor(x)
        return vals[lo] + (vals[min(lo + 1, n - 1)] - vals[lo]) * (x - lo)

    c25, c50, c75 = cuantil(.25), cuantil(.5), cuantil(.75)

    def cuartil(v):
        return 4 if v <= c25 else 3 if v <= c50 else 2 if v <= c75 else 1

    efw = {i: {'efw_avg': round(prom[i]['s'], 3), 'efw_quartile': cuartil(prom[i]['s']), 'efw_rank': p + 1,
               'efw_n_years': prom[i]['n'], 'region': reg[i], 'income_class': ing[i],
               **{f'efw_a{j}': (None if prom[i][f'a{j}'] is None else round(prom[i][f'a{j}'], 3)) for j in range(1, 6)}}
           for p, i in enumerate(orden)}
    return out, meta, efw, (anios[0], anios[-1]), (c25, c50, c75)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    salida = Path(sys.argv[sys.argv.index('--salida') + 1]) if '--salida' in sys.argv else DATA
    if '--salida' in sys.argv:
        args = [a for a in args if Path(a) != salida]
    if args:
        xlsx = Path(args[0])
    else:
        xlsx = DATA / f'efotw-{INFORME}-master-index-data-for-researchers-iso.xlsx'
        if not xlsx.exists():
            print(f'Descargando el archivo maestro del informe {INFORME}…')
            req = urllib.request.Request(URL.format(INFORME), headers={'User-Agent': 'Mozilla/5.0'})
            xlsx.write_bytes(urllib.request.urlopen(req, timeout=120).read())
    m = re.search(r'efotw-(\d{4})', xlsx.name)
    informe = int(m.group(1)) if m else None
    d = leer(xlsx)
    desfase, n_areas, ult = corregir_desfase(d)
    if informe and informe != ult + 2:
        alto(f'el archivo dice informe {informe} pero el último dato es {ult} (se esperaba {ult + 2})')
    panel, meta, efw, (a0, a1), cortes = construir(d)

    # consolidated.json: se actualizan los campos del índice, el resto queda como estaba
    base = json.loads((DATA / 'consolidated.json').read_text(encoding='utf-8'))
    por_iso = {r['iso']: r for r in base}
    if set(por_iso) != set(efw):
        alto(f'consolidated.json y el índice no tienen las mismas jurisdicciones: '
             f'{sorted(set(por_iso) ^ set(efw))}')
    for iso, campos in efw.items():
        por_iso[iso]['country'] = d['nombre'][iso]
        por_iso[iso].update(campos)
    cons = sorted(por_iso.values(), key=lambda r: r['efw_rank'])

    salida.mkdir(parents=True, exist_ok=True)
    (salida / 'efw_panel_map.json').write_text(json.dumps(panel), encoding='utf-8')
    (salida / 'efw_country_meta.json').write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding='utf-8')
    (salida / 'consolidated.json').write_text(json.dumps(cons, indent=2, ensure_ascii=False), encoding='utf-8')
    comp = {'fuente': f'Fraser Institute, Economic Freedom of the World: {informe} Annual Report, hoja «{d["hojas"][0]}»',
            'nota': 'Bolivia, hoja oficial del informe (sin encadenar: puede diferir del panel en los años viejos). '
                    'Cada componente de 0 a 10; «· dato» es su valor de origen cuando la hoja lo trae (la inflación '
                    'en %, el crecimiento del dinero…); «Rank» es el puesto oficial en esa área.',
            'anios': d['componentes_bol']}
    (salida / 'efw_componentes_bol.json').write_text(json.dumps(comp, indent=1, ensure_ascii=False), encoding='utf-8')

    b = panel[str(ult)]['BOL']
    pb = 1 + sum(1 for v in panel[str(ult)].values() if v['s'] > b['s'])
    print(f'✓ {xlsx.name}: informe {informe}, datos {min(map(int, panel))}–{ult}, {len(panel[str(ult)])} jurisdicciones en {ult}')
    print(f'  desfase del panel frente a la hoja oficial en {ult}: {desfase:+.2f} en {n_areas} áreas'
          + (' → descontado en todo el panel' if desfase else ' (ninguno)'))
    print(f'  promedio {a0}–{a1}; cortes de cuartil {cortes[0]:.3f} / {cortes[1]:.3f} / {cortes[2]:.3f}')
    print(f'  Bolivia {ult}: {b["s"]:.2f} (puesto {pb} de {len(panel[str(ult)])}) · áreas '
          + ' · '.join(f'{b[k]:.2f}' for k in K[1:]))
    print(f'  escrito en {salida}')


if __name__ == '__main__':
    main()
